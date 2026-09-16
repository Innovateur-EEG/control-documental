import os
from docx import Document

def crear_plantilla():
    # Asegurarnos de que el directorio exista
    os.makedirs('app/templates', exist_ok=True)
    
    doc = Document()
    doc.add_heading('FORMATO ÚNICO DE REGISTRO', 0)
    
    doc.add_paragraph('El presente documento certifica el registro y verificación de datos en la plataforma de control documental.')
    
    doc.add_heading('1. Datos Generales', level=1)
    doc.add_paragraph('Nombre / Razón Social: {{ nombre }}')
    doc.add_paragraph('Registro Federal de Contribuyentes (RFC): {{ rfc }}')
    doc.add_paragraph('Dirección Registrada: {{ direccion }}')
    doc.add_paragraph('Teléfono de Contacto: {{ telefono }}')
    
    doc.add_heading('2. Datos Específicos', level=1)
    # Sintaxis Jinja2 soportada nativamente por docxtpl dentro de Word
    doc.add_paragraph('{% if fecha_const %}Fecha de Constitución: {{ fecha_const }}')
    doc.add_paragraph('Notario Público: {{ notario }}{% endif %}')
    
    doc.add_paragraph('{% if curp %}Clave Única de Registro de Población (CURP): {{ curp }}')
    doc.add_paragraph('Estado Civil: {{ estado_civil }}{% endif %}')
    
    doc.add_paragraph('\n\n__________________________________\nFirma de Conformidad')
    
    ruta_salida = 'app/templates/formato_base.docx'
    doc.save(ruta_salida)
    print(f"Plantilla maestra creada exitosamente en: {ruta_salida}")

if __name__ == "__main__":
    crear_plantilla()