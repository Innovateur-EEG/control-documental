"""Regression tests for Excel Jinja template rendering."""

import io

import openpyxl

from app.utils.documentos import procesar_excel


def _workbook_bytes(value):
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = value
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def _rendered_cell(value, context):
    result = procesar_excel(_workbook_bytes(value), context)
    workbook = openpyxl.load_workbook(io.BytesIO(result), data_only=False)
    return workbook.active["A1"].value


def test_excel_renders_variables_with_optional_whitespace():
    assert _rendered_cell("{{ nombre }} / {{rfc}}", {"nombre": "Ana", "rfc": "ABC"}) == "Ana / ABC"


def test_excel_renders_registered_jinja_filters():
    assert _rendered_cell("{{ genero | mark(\"Masculino\") }}", {"genero": "Masculino"}) == "X"
    assert _rendered_cell("{{ interior | prefix(\", interior: \") }}", {"interior": None}) == ""


def test_excel_keeps_plain_text_and_formula_cells_unchanged():
    workbook = openpyxl.Workbook()
    workbook.active["A1"] = "Texto fijo"
    workbook.active["A2"] = "=1+1"
    output = io.BytesIO()
    workbook.save(output)
    result = procesar_excel(output.getvalue(), {})
    rendered = openpyxl.load_workbook(io.BytesIO(result), data_only=False)
    assert rendered.active["A1"].value == "Texto fijo"
    assert rendered.active["A2"].value == "=1+1"
