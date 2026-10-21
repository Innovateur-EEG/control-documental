"""Regression tests for Word template rendering."""
from io import BytesIO

from docx import Document

from app.utils.documentos import procesar_excel, procesar_word


def test_procesar_word_renders_docx_with_custom_jinja_filters():
    source = Document()
    source.add_paragraph("Hola {{ nombre }} {{ estado | mark('Activo') }}")
    input_bytes = BytesIO()
    source.save(input_bytes)

    rendered_bytes = procesar_word(
        input_bytes.getvalue(),
        {"nombre": "Natalia", "estado": "Activo"},
    )
    rendered = Document(BytesIO(rendered_bytes))
    text = "\n".join(paragraph.text for paragraph in rendered.paragraphs)

    assert "Hola Natalia X" in text

def test_procesar_excel_renders_saved_capture_values_and_preserves_workbook_structure():
    from openpyxl import Workbook, load_workbook

    source = Workbook()
    sheet = source.active
    sheet.title = "Captura"
    sheet["A1"] = "Nombre: {{ nombre }}"
    sheet["B2"] = "{{ estado }}"
    sheet["C3"] = 42
    sheet["D4"] = "=1+1"
    sheet["A1"].number_format = "@"
    extra = source.create_sheet("Anexo")
    extra["A1"] = "Valor: {{ importe }}"

    input_bytes = BytesIO()
    source.save(input_bytes)

    rendered_bytes = procesar_excel(
        input_bytes.getvalue(),
        {"nombre": "Persona de prueba", "estado": "Activo", "importe": "125.50"},
    )
    rendered = load_workbook(BytesIO(rendered_bytes), data_only=False)

    assert rendered.sheetnames == ["Captura", "Anexo"]
    assert rendered["Captura"]["A1"].value == "Nombre: Persona de prueba"
    assert rendered["Captura"]["B2"].value == "Activo"
    assert rendered["Captura"]["C3"].value == 42
    assert rendered["Captura"]["D4"].value == "=1+1"
    assert rendered["Captura"]["A1"].number_format == "@"
    assert rendered["Anexo"]["A1"].value == "Valor: 125.50"

def test_procesar_word_renders_capture_values_inside_template_tables():
    source = Document()
    table = source.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Representante"
    table.cell(0, 1).text = "{{ representante }}"
    input_bytes = BytesIO()
    source.save(input_bytes)

    rendered_bytes = procesar_word(
        input_bytes.getvalue(),
        {"representante": "Nombre guardado"},
    )
    rendered = Document(BytesIO(rendered_bytes))

    assert rendered.tables[0].cell(0, 0).text == "Representante"
    assert rendered.tables[0].cell(0, 1).text == "Nombre guardado"


def test_template_batch_fails_closed_instead_of_returning_partial_package(monkeypatch):
    from types import SimpleNamespace
    import app.utils.documentos as documentos

    monkeypatch.setattr(documentos, "procesar_word", lambda data, context: b"rendered")
    monkeypatch.setattr(documentos, "procesar_excel", lambda data, context: b"rendered")
    monkeypatch.setattr(documentos, "convertir_a_pdf", lambda data, extension: b"pdf")
    templates = [
        SimpleNamespace(nombre="Correcta", ruta_plantilla_word="ok.docx"),
        SimpleNamespace(nombre="Fallida", ruta_plantilla_word="bad.docx"),
    ]

    def download(path):
        if path == "bad.docx":
            raise RuntimeError("storage unavailable")
        return b"template"

    pdfs, errors = documentos.procesar_plantillas_a_pdf(templates, {}, download)

    assert pdfs == []
    assert len(errors) == 1
    assert "Fallida" in errors[0]
    assert "storage unavailable" in errors[0]


def test_template_batch_rejects_missing_and_unsupported_templates():
    from types import SimpleNamespace
    from app.utils.documentos import procesar_plantillas_a_pdf

    templates = [
        SimpleNamespace(nombre="Sin archivo", ruta_plantilla_word=None),
        SimpleNamespace(nombre="Formato antiguo", ruta_plantilla_word="archivo.doc"),
    ]
    pdfs, errors = procesar_plantillas_a_pdf(templates, {}, lambda path: b"unused")

    assert pdfs == []
    assert len(errors) == 2
    assert "Sin archivo" in errors[0]
    assert "Formato antiguo" in errors[1]

def test_template_batch_returns_all_pdfs_when_every_template_succeeds(monkeypatch):
    from types import SimpleNamespace
    import app.utils.documentos as documentos

    calls = []
    monkeypatch.setattr(documentos, "procesar_word", lambda data, context: b"word-rendered")
    monkeypatch.setattr(documentos, "procesar_excel", lambda data, context: b"excel-rendered")
    monkeypatch.setattr(
        documentos,
        "convertir_a_pdf",
        lambda data, extension: calls.append((data, extension)) or b"pdf-" + extension.encode(),
    )
    templates = [
        SimpleNamespace(nombre="Word", ruta_plantilla_word="one.docx"),
        SimpleNamespace(nombre="Excel", ruta_plantilla_word="two.xlsx"),
    ]

    pdfs, errors = documentos.procesar_plantillas_a_pdf(
        templates, {"nombre": "Prueba"}, lambda path: path.encode()
    )

    assert pdfs == [b"pdf-.docx", b"pdf-.xlsx"]
    assert errors == []
    assert calls == [(b"word-rendered", ".docx"), (b"excel-rendered", ".xlsx")]


def test_template_batch_discards_prior_pdfs_when_conversion_fails(monkeypatch):
    from types import SimpleNamespace
    import app.utils.documentos as documentos

    monkeypatch.setattr(documentos, "procesar_word", lambda data, context: b"rendered")

    conversion_count = 0

    def convert(data, extension):
        nonlocal conversion_count
        conversion_count += 1
        if conversion_count == 2:
            raise RuntimeError("conversion failed")
        return b"pdf"

    monkeypatch.setattr(documentos, "convertir_a_pdf", convert)
    templates = [
        SimpleNamespace(nombre="Primera", ruta_plantilla_word="first.docx"),
        SimpleNamespace(nombre="Segunda", ruta_plantilla_word="second.docx"),
    ]

    pdfs, errors = documentos.procesar_plantillas_a_pdf(
        templates,
        {},
        lambda path: path.encode(),
    )

    assert pdfs == []
    assert len(errors) == 1
    assert "Segunda" in errors[0]
    assert "conversion failed" in errors[0]

def test_generated_package_path_never_falls_back_to_legacy_signed_file():
    from app.utils.documentos import obtener_ruta_paquete_generado

    assert obtener_ruta_paquete_generado(
        {"ruta_paquete_generado": "generated/current.pdf"},
        "signed/legacy.pdf",
        "Firmado y Cargado",
    ) == "generated/current.pdf"
    assert obtener_ruta_paquete_generado(
        {},
        "signed/legacy.pdf",
        "Firmado y Cargado",
    ) is None
    assert obtener_ruta_paquete_generado(
        {},
        "generated/legacy.pdf",
        "Generado",
    ) == "generated/legacy.pdf"

